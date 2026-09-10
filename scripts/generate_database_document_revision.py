from __future__ import annotations

import os
from pathlib import Path
from shutil import copyfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path("/tmp/tongji-dbdoc-revision/数据库设计文档.docx")
OUTPUT = ROOT / "document/数据库设计文档_修订稿.docx"
ASSET_DIR = Path("/tmp/tongji-dbdoc-revision/revision-assets")
# 字体路径不写死个人目录：可用 DBDOC_FONT 覆盖，默认取当前用户的 macOS 字体目录。
FONT_PATH = Path(
    os.environ.get("DBDOC_FONT", str(Path.home() / "Library/Fonts/NotoSerifSC-Variable.ttf"))
)


TABLES = [
    ("DEPARTMENTS_TRAINING", "部门培训预算表", "记录 2026 年部门培训预算、已用预算和由数据库计算的余额。", [
        ("DEPT_ID", "NUMBER", "PK，IDENTITY", "部门编号"),
        ("DEPT_NAME", "VARCHAR2(100)", "NOT NULL，UNIQUE", "部门名称"),
        ("BUDGET_YEAR", "NUMBER(4)", "NOT NULL，DEFAULT 2026，CHECK=2026", "预算年度"),
        ("ANNUAL_BUDGET", "NUMBER(12,2)", "NOT NULL，DEFAULT 0，>=0", "年度预算"),
        ("USED_BUDGET", "NUMBER(12,2)", "NOT NULL，DEFAULT 0，0~年度预算", "已使用预算"),
        ("REMAIN_BUDGET", "NUMBER(12,2)", "虚拟列", "ANNUAL_BUDGET-USED_BUDGET"),
        ("CREATED_AT / UPDATED_AT", "TIMESTAMP(0)", "NOT NULL", "创建、更新时间"),
    ]),
    ("EMPLOYEES", "员工表", "记录登录主体、组织归属和员工在职状态，是流程数据的基础实体。", [
        ("EMP_ID", "NUMBER", "PK，IDENTITY", "员工编号"),
        ("LOGIN_NAME", "VARCHAR2(50)", "NOT NULL，UNIQUE，小写约束", "登录名"),
        ("PASSWORD_HASH", "VARCHAR2(255)", "NOT NULL", "BCrypt 密码哈希"),
        ("EMP_NAME", "VARCHAR2(50)", "NOT NULL", "员工姓名"),
        ("DEPT_ID", "NUMBER", "NOT NULL，FK->部门", "所属部门"),
        ("MANAGER_EMP_ID", "NUMBER", "FK->EMPLOYEES，可空", "直属主管扩展字段"),
        ("POSITION / EMAIL / PHONE", "VARCHAR2", "EMAIL UNIQUE，PHONE(20)", "职位、邮箱、电话"),
        ("HIRE_DATE", "DATE", "可空", "入职日期"),
        ("STATUS", "VARCHAR2(20)", "ACTIVE / RESIGNED", "在职状态"),
        ("CREATED_AT / UPDATED_AT", "TIMESTAMP(0)", "NOT NULL", "创建、更新时间"),
    ]),
    ("ROLES", "系统角色表", "保存固定四类角色的稳定编码和显示名称，权限策略由后端维护。", [
        ("ROLE_ID", "NUMBER", "PK，IDENTITY", "角色编号"),
        ("ROLE_CODE", "VARCHAR2(30)", "NOT NULL，UNIQUE", "EMPLOYEE、DEPT_MANAGER、HR、ADMIN"),
        ("ROLE_NAME", "VARCHAR2(50)", "NOT NULL，UNIQUE", "角色名称"),
        ("CREATED_AT", "TIMESTAMP(0)", "NOT NULL", "创建时间"),
    ]),
    ("USER_ROLES", "员工角色关联表", "实现员工与角色的多对多授权关系。", [
        ("USER_ROLE_ID", "NUMBER", "PK，IDENTITY", "关联编号"),
        ("EMP_ID", "NUMBER", "NOT NULL，FK->员工", "员工编号"),
        ("ROLE_ID", "NUMBER", "NOT NULL，FK->角色", "角色编号"),
        ("CREATED_AT", "TIMESTAMP(0)", "NOT NULL", "创建时间"),
        ("(EMP_ID, ROLE_ID)", "-", "UNIQUE", "禁止重复授予同一角色"),
    ]),
    ("BLACKLIST", "培训黑名单表", "限制处于生效期的员工创建新报名，历史解除记录予以保留。", [
        ("BLACK_ID", "NUMBER", "PK，IDENTITY", "黑名单编号"),
        ("EMP_ID", "NUMBER", "NOT NULL，FK->员工", "员工编号"),
        ("REASON", "VARCHAR2(500)", "NOT NULL", "列入原因"),
        ("START_AT / END_AT", "TIMESTAMP(0)", "END_AT>=START_AT", "生效、解除时间"),
        ("STATUS", "VARCHAR2(20)", "ACTIVE / RELEASED", "黑名单状态"),
        ("CREATED_AT / UPDATED_AT", "TIMESTAMP(0)", "NOT NULL", "创建、更新时间"),
    ]),
    ("TRAINERS", "讲师表", "记录课程主讲师基础信息和星级。", [
        ("TRAINER_ID", "NUMBER", "PK，IDENTITY", "讲师编号"),
        ("TRAINER_NAME", "VARCHAR2(50)", "NOT NULL", "讲师姓名"),
        ("TITLE / COMPANY", "VARCHAR2", "可空", "职称、所属公司"),
        ("PHONE / EMAIL", "VARCHAR2", "PHONE(20)", "联系方式"),
        ("STAR_LEVEL", "NUMBER(2,1)", "DEFAULT 3，1~5", "讲师星级"),
        ("IS_INTERNAL", "VARCHAR2(1)", "DEFAULT Y，Y/N", "是否内部讲师"),
        ("CREATED_AT / UPDATED_AT", "TIMESTAMP(0)", "NOT NULL", "创建、更新时间"),
    ]),
    ("TRAINING_COURSES", "培训课程表", "记录课程、单一主讲师、主办部门、预算和发布条件。", [
        ("COURSE_ID", "NUMBER", "PK，IDENTITY", "课程编号"),
        ("COURSE_NAME / COURSE_TYPE", "VARCHAR2", "均 NOT NULL", "课程名称、类型"),
        ("DURATION_HOURS", "NUMBER(4,1)", ">0", "课程学时"),
        ("TRAINER_ID", "NUMBER", "FK->讲师，可空草稿", "唯一主讲师"),
        ("MAX_STUDENTS", "NUMBER(6)", ">0", "最大人数"),
        ("START_AT / END_AT / LOCATION", "TIMESTAMP / VARCHAR2", "发布时必填，END>START", "培训时间、地点"),
        ("COURSE_STATUS", "VARCHAR2(20)", "DRAFT / PUBLISHED / CLOSED", "课程状态"),
        ("BUDGET_AMOUNT", "NUMBER(12,2)", "DEFAULT 0，>=0", "预算金额"),
        ("DEPT_ID", "NUMBER", "FK->部门，可空草稿", "主办部门"),
        ("PRE_TEST_URL / POST_TEST_URL / MATERIAL_URL", "VARCHAR2(500)", "可空", "测试、课件链接"),
        ("CREATED_AT / UPDATED_AT", "TIMESTAMP(0)", "NOT NULL", "创建、更新时间"),
    ]),
    ("TRAINING_REQUESTS", "培训申请表", "记录员工申请、部门审批与 HR 备案的完整留痕。", [
        ("REQUEST_ID", "NUMBER", "PK，IDENTITY", "申请编号"),
        ("EMP_ID / COURSE_ID", "NUMBER", "NOT NULL，FK", "申请员工、目标课程"),
        ("REASON", "VARCHAR2(500)", "NOT NULL", "申请理由"),
        ("STATUS", "VARCHAR2(20)", "PENDING / DEPT_APPROVED / DEPT_REJECTED / HR_FILED", "审批状态"),
        ("DEPT_APPROVER_EMP_ID", "NUMBER", "FK->员工，可空", "部门审批人"),
        ("DEPT_APPROVED_AT / DEPT_APPROVE_REMARK", "TIMESTAMP / VARCHAR2(500)", "可空", "部门审批留痕"),
        ("HR_FILER_EMP_ID", "NUMBER", "FK->员工，可空", "HR 备案人"),
        ("HR_FILED_AT / HR_FILE_REMARK", "TIMESTAMP / VARCHAR2(500)", "备案时必填", "HR 备案留痕"),
        ("REQUESTED_AT / UPDATED_AT", "TIMESTAMP(0)", "NOT NULL", "提交、更新时间"),
    ]),
    ("TRAINING_REGISTRATIONS", "培训报名表", "报名必须引用同员工、同课程且已经 HR 备案的申请。", [
        ("REG_ID", "NUMBER", "PK，IDENTITY", "报名编号"),
        ("REQUEST_ID / EMP_ID / COURSE_ID", "NUMBER", "NOT NULL，复合 FK->申请", "可追溯报名来源"),
        ("STATUS", "VARCHAR2(20)", "REGISTERED / SIGNED_IN / COMPLETED / ABSENT / CANCELED", "报名状态"),
        ("REGISTERED_AT", "TIMESTAMP(0)", "NOT NULL", "报名时间"),
        ("COMPLETED_AT / ACTUAL_HOURS", "TIMESTAMP / NUMBER(4,1)", "完成时均非空", "完成与实际学时"),
        ("CANCELED_AT / CANCEL_REASON", "TIMESTAMP / VARCHAR2(500)", "取消时记录", "取消留痕"),
        ("UPDATED_AT", "TIMESTAMP(0)", "NOT NULL", "更新时间"),
    ]),
    ("TRAINING_ATTENDANCE", "培训签到记录表", "保存签到事实；一条报名至多对应一条签到记录。", [
        ("ATTEND_ID", "NUMBER", "PK，IDENTITY", "签到记录编号"),
        ("REG_ID", "NUMBER", "NOT NULL，UNIQUE，FK->报名", "报名编号"),
        ("SIGNIN_TYPE", "VARCHAR2(20)", "SCAN / MANUAL", "扫码或补签"),
        ("SIGNED_IN_AT", "TIMESTAMP(0)", "NOT NULL", "实际签到时间"),
        ("LATENESS_MINUTES", "NUMBER", "DEFAULT 0，>=0", "迟到分钟数"),
        ("DEDUCT_HOURS", "NUMBER(4,1)", "DEFAULT 0，>=0", "迟到扣减学时"),
        ("REMARK", "VARCHAR2(500)", "可空；补签必填", "备注"),
        ("CREATED_AT", "TIMESTAMP(0)", "NOT NULL", "创建时间"),
    ]),
    ("TRAINER_RATINGS", "讲师评分表", "记录员工对课程实际主讲师的评分和 HR 复核。", [
        ("RATING_ID", "NUMBER", "PK，IDENTITY", "评分编号"),
        ("COURSE_ID / TRAINER_ID", "NUMBER", "复合 FK->课程主讲师", "课程、实际讲师"),
        ("EMP_ID", "NUMBER", "NOT NULL，FK->员工", "评分员工"),
        ("SCORE / RATING_COMMENT", "NUMBER(2,1) / VARCHAR2(500)", "SCORE 1~5", "评分、评价"),
        ("HR_VERIFIED", "VARCHAR2(1)", "DEFAULT N，Y/N", "是否复核"),
        ("HR_VERIFIER_EMP_ID / VERIFIED_AT / HR_COMMENT", "NUMBER / TIMESTAMP / VARCHAR2(500)", "复核时前两项必填", "复核留痕"),
        ("RATED_AT", "TIMESTAMP(0)", "NOT NULL", "评分时间"),
        ("(EMP_ID, COURSE_ID, TRAINER_ID)", "-", "UNIQUE", "每员工每课程主讲师仅一次评分"),
    ]),
    ("TRAINING_TESTS", "培训效果测试表", "记录训前、训后测试及其录入人。", [
        ("TEST_ID", "NUMBER", "PK，IDENTITY", "测试编号"),
        ("EMP_ID / COURSE_ID", "NUMBER", "NOT NULL，FK", "参加员工、课程"),
        ("TEST_TYPE", "VARCHAR2(10)", "PRE / POST", "测试类型"),
        ("SCORE", "NUMBER(5,2)", "0~100", "测试成绩"),
        ("RECORDED_BY_EMP_ID", "NUMBER", "NOT NULL，FK->员工", "录入人"),
        ("TESTED_AT", "TIMESTAMP(0)", "NOT NULL", "记录时间"),
        ("(EMP_ID, COURSE_ID, TEST_TYPE)", "-", "UNIQUE", "同类测试仅一条"),
    ]),
    ("TRAINING_CERTIFICATES", "培训证书表", "记录培训完成后生成的唯一证书及发证、通知留痕。", [
        ("CERT_ID", "NUMBER", "PK，IDENTITY", "证书记录编号"),
        ("EMP_ID / COURSE_ID", "NUMBER", "NOT NULL，FK", "获证员工、课程"),
        ("CERT_CODE", "VARCHAR2(80)", "NOT NULL，UNIQUE", "唯一证书编号"),
        ("ISSUE_DATE / EXPIRE_DATE", "DATE", "EXPIRE>=ISSUE", "发证、到期日期"),
        ("NOTIFIED / NOTIFIED_AT", "VARCHAR2(1) / TIMESTAMP", "Y/N；Y 时有通知时间", "通知状态"),
        ("ISSUED_BY_EMP_ID", "NUMBER", "NOT NULL，FK->员工", "发证人"),
        ("CREATED_AT", "TIMESTAMP(0)", "NOT NULL", "创建时间"),
        ("(EMP_ID, COURSE_ID)", "-", "UNIQUE", "同员工同课程仅一证"),
    ]),
]


def set_east_asia(run, font_name="宋体", size=None, bold=None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def paragraph(doc, text="", style="文档正文", align=None, first_line=True):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if first_line:
        p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.35
    run = p.add_run(text)
    set_east_asia(run, "宋体", 10.5)
    return p


def heading(doc, text, level=1):
    style = f"heading {level}"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_east_asia(run, "黑体", {1: 15, 2: 14, 3: 12, 4: 11}[level], True)
    return p


def remove_legacy_heading_numbering(doc):
    for level in range(1, 5):
        style = doc.styles[f"heading {level}"]
        p_pr = style.element.get_or_add_pPr()
        num_pr = p_pr.find(qn("w:numPr"))
        if num_pr is not None:
            p_pr.remove(num_pr)


def caption(doc, text):
    p = doc.add_paragraph(style="题注")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_east_asia(run, "宋体", 10.5)
    return p


def add_table(doc, headers, rows, widths=(2.7, 2.7, 5.3, 5.2)):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    header_cells = table.rows[0].cells
    for cell, text, width in zip(header_cells, headers, widths):
        cell.width = Cm(width)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cell.paragraphs[0].add_run(text)
        set_east_asia(run, "宋体", 9, True)
    for row in rows:
        cells = table.add_row().cells
        for cell, text, width in zip(cells, row, widths):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(text))
            set_east_asia(run, "宋体", 8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def arrow(draw, a, b, label="", font=None):
    draw.line((a, b), fill="#222222", width=3)
    x1, y1 = a
    x2, y2 = b
    dx, dy = x2 - x1, y2 - y1
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    tip = (x2, y2)
    left = (x2 - 16 * ux + 8 * px, y2 - 16 * uy + 8 * py)
    right = (x2 - 16 * ux - 8 * px, y2 - 16 * uy - 8 * py)
    draw.polygon([tip, left, right], fill="#222222")
    if label:
        mx, my = (x1 + x2) / 2, (y1 + y2) / 2
        bbox = draw.textbbox((0, 0), label, font=font)
        draw.rounded_rectangle((mx - (bbox[2]-bbox[0])/2 - 8, my - 20, mx + (bbox[2]-bbox[0])/2 + 8, my + 20), 5, fill="white")
        draw.text((mx - (bbox[2]-bbox[0])/2, my - (bbox[3]-bbox[1])/2), label, font=font, fill="#222222")


def entity_diagram(path, title, nodes, relations, size=(1800, 1120)):
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(str(FONT_PATH), 46)
    body_font = ImageFont.truetype(str(FONT_PATH), 30)
    small_font = ImageFont.truetype(str(FONT_PATH), 26)
    draw.text((60, 32), title, font=title_font, fill="#111111")
    for start, end, label in relations:
        sx, sy, sw, sh = nodes[start]
        ex, ey, ew, eh = nodes[end]
        a = (sx + sw / 2, sy + sh / 2)
        b = (ex + ew / 2, ey + eh / 2)
        arrow(draw, a, b, label, small_font)
    for name, (x, y, w, h) in nodes.items():
        draw.rounded_rectangle((x, y, x + w, y + h), 12, fill="#fafafa", outline="#202020", width=3)
        bbox = draw.textbbox((0, 0), name, font=body_font)
        draw.text((x + (w - bbox[2]) / 2, y + (h - bbox[3]) / 2 - 2), name, font=body_font, fill="#111111")
    draw.text((60, size[1] - 58), "关系标注采用 1、N、0..1；以当前 Oracle DDL 与后端流程为准。", font=small_font, fill="#333333")
    image.save(path)


def build_diagrams():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    overall = {
        "部门": (90, 210, 210, 80), "员工": (420, 150, 210, 80), "角色": (420, 420, 210, 80),
        "员工角色": (740, 420, 210, 80), "黑名单": (740, 130, 210, 80), "讲师": (90, 700, 210, 80),
        "课程": (420, 700, 210, 80), "申请": (740, 700, 210, 80), "报名": (1060, 700, 210, 80),
        "签到": (1390, 590, 210, 80), "评分": (1390, 750, 210, 80), "测试": (1060, 890, 210, 80), "证书": (1390, 920, 210, 80),
    }
    entity_diagram(ASSET_DIR / "overall.png", "企业内部培训管理系统总体 E-R 图", overall, [
        ("部门", "员工", "1:N"), ("部门", "课程", "1:N"), ("员工", "员工角色", "1:N"), ("角色", "员工角色", "1:N"),
        ("员工", "黑名单", "1:N"), ("讲师", "课程", "1:N"), ("员工", "申请", "1:N"), ("课程", "申请", "1:N"),
        ("申请", "报名", "1:N 历史"), ("报名", "签到", "1:0..1"), ("课程", "评分", "1:N"), ("员工", "评分", "1:N"),
        ("课程", "测试", "1:N"), ("员工", "测试", "1:N"), ("课程", "证书", "1:N"), ("员工", "证书", "1:N"),
    ])
    base = {"部门": (100, 350, 240, 90), "员工": (500, 190, 240, 90), "角色": (500, 530, 240, 90), "员工角色": (900, 530, 240, 90), "黑名单": (900, 140, 240, 90), "讲师": (100, 690, 240, 90), "课程": (500, 780, 240, 90)}
    entity_diagram(ASSET_DIR / "basic.png", "组织与培训基础模块 E-R 图", base, [("部门", "员工", "1:N"), ("部门", "课程", "1:N"), ("员工", "员工角色", "1:N"), ("角色", "员工角色", "1:N"), ("员工", "黑名单", "1:N"), ("讲师", "课程", "1:N")], (1400, 1000))
    flow = {"员工": (80, 400, 220, 90), "课程": (430, 170, 220, 90), "申请": (770, 400, 220, 90), "报名": (1110, 400, 220, 90)}
    entity_diagram(ASSET_DIR / "request.png", "需求与报名模块 E-R 图", flow, [("员工", "申请", "1:N"), ("课程", "申请", "1:N"), ("申请", "报名", "HR_FILED 后")], (1400, 850))
    result = {"员工": (80, 410, 210, 80), "课程": (410, 160, 210, 80), "报名": (410, 650, 210, 80), "签到": (750, 650, 210, 80), "评分": (750, 250, 210, 80), "测试": (1080, 400, 210, 80), "证书": (1080, 650, 210, 80), "讲师": (410, 400, 210, 80)}
    entity_diagram(ASSET_DIR / "result.png", "培训执行与成果模块 E-R 图", result, [("课程", "报名", "1:N"), ("报名", "签到", "1:0..1"), ("员工", "评分", "1:N"), ("课程", "评分", "1:N"), ("讲师", "评分", "1:N"), ("员工", "测试", "1:N"), ("课程", "测试", "1:N"), ("员工", "证书", "1:N"), ("课程", "证书", "1:N")], (1400, 900))
    physical = {"部门": (60, 150, 180, 65), "员工": (340, 120, 180, 65), "角色": (340, 280, 180, 65), "员工角色": (630, 280, 180, 65), "黑名单": (630, 110, 180, 65), "讲师": (60, 610, 180, 65), "课程": (340, 610, 180, 65), "申请": (630, 610, 180, 65), "报名": (920, 610, 180, 65), "签到": (1210, 480, 180, 65), "评分": (1210, 650, 180, 65), "测试": (920, 790, 180, 65), "证书": (1210, 820, 180, 65)}
    entity_diagram(ASSET_DIR / "physical.png", "数据库关系图（主键、外键关系）", physical, [("部门", "员工", "DEPT_ID"), ("部门", "课程", "DEPT_ID"), ("员工", "员工角色", "EMP_ID"), ("角色", "员工角色", "ROLE_ID"), ("员工", "黑名单", "EMP_ID"), ("讲师", "课程", "TRAINER_ID"), ("员工", "申请", "EMP_ID"), ("课程", "申请", "COURSE_ID"), ("申请", "报名", "复合 FK"), ("报名", "签到", "REG_ID UNIQUE"), ("课程", "评分", "课程+讲师"), ("员工", "评分", "EMP_ID"), ("课程", "测试", "COURSE_ID"), ("员工", "测试", "EMP_ID"), ("课程", "证书", "COURSE_ID"), ("员工", "证书", "EMP_ID")], (1500, 1000))


def clear_after_cover(doc):
    body = doc._body._element
    # Direct body children 0..20 hold the cover title and member list. The
    # following empty paragraphs and old section break create a blank page
    # before the regenerated table of contents.
    for index, child in reversed(list(enumerate(body))):
        if index >= 21 and child.tag != qn("w:sectPr"):
            body.remove(child)


def add_toc(doc):
    p = doc.add_paragraph(style="封面副标题")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.page_break_before = True
    r = p.add_run("目  录")
    set_east_asia(r, "黑体", 18, True)
    entries = [
        "1. 企业内部培训管理系统数据需求", "1.1 组织与培训基础数据需求", "1.2 需求与报名功能数据需求", "1.3 培训执行与成果数据需求",
        "2. 概念设计", "2.1 总体 E-R 图", "2.2 组织与培训基础模块 E-R 图", "2.3 需求与报名模块 E-R 图", "2.4 培训执行与成果模块 E-R 图",
        "3. 逻辑设计", "3.1 表的设计", "3.2 完整性、状态与索引设计", "3.3 数据库关系图", "附录A. 图表索引"
    ]
    for entry in entries:
        p = doc.add_paragraph(style="toc 1")
        p.paragraph_format.line_spacing = 1.25
        r = p.add_run(entry)
        set_east_asia(r, "宋体", 10.5)


def add_content(doc):
    doc.add_page_break()
    heading(doc, "1. 企业内部培训管理系统数据需求", 1)
    paragraph(doc, "企业内部培训管理系统的数据需求围绕企业培训全流程管理展开，覆盖组织人员、培训资源、预算控制、申请审批、报名签到、效果评估和证书管理。系统以 Oracle 关系数据库保存业务事实，通过主键、外键、唯一约束和检查约束保证数据一致性。")
    heading(doc, "1.1 组织与培训基础数据需求", 2)
    paragraph(doc, "核心功能：维护部门年度培训预算、员工登录身份与组织归属、角色授权、黑名单、讲师和培训课程，为后续业务流程提供可追溯的基础数据。")
    for text in ["员工信息：记录登录名、密码哈希、部门、职位、联系方式和在职状态；部门名称不作为外键使用。", "角色信息：固定员工、部门主管、HR、管理员四类角色，通过员工角色关联表实现多角色授权。", "培训课程：一门课程在本期只配置一名主讲师；课程发布前必须具备主讲师、部门、时间、地点、人数、学时和预算。", "部门培训预算：余额由年度预算减已用预算自动计算，避免页面或服务端重复维护余额。"]:
        paragraph(doc, "·" + text, first_line=False)
    heading(doc, "1.2 需求与报名功能数据需求", 2)
    paragraph(doc, "核心功能：管理员工提交培训申请、部门主管审批、HR 备案及报名。员工不能绕过申请直接报名，报名必须引用同一员工、同一课程且状态为 HR_FILED 的申请。")
    for text in ["培训申请：记录申请理由、部门审批人及意见、HR 备案人及意见，并以状态机保存流程进度。", "培训报名：记录报名、签到、完成、缺勤、取消状态。取消记录保留；同一员工同一课程可在取消后重新报名，但任意时刻只能有一条有效报名。"]:
        paragraph(doc, "·" + text, first_line=False)
    heading(doc, "1.3 培训执行与成果数据需求", 2)
    paragraph(doc, "核心功能：记录一次培训的签到事实、实际学时、讲师评分、训前训后测试和证书发放，形成培训结果闭环。")
    for text in ["签到记录：一条报名至多一条签到，扫码与补签均记录实际签到时间、迟到分钟数、扣减学时及备注。", "讲师评分：评分必须对应课程实际主讲师；同一员工对同一课程主讲师只能评分一次。", "培训测试与证书：同一员工同一课程的 PRE、POST 测试各至多一条；证书在满足完成条件后生成，员工与课程组合唯一。"]:
        paragraph(doc, "·" + text, first_line=False)

    heading(doc, "2. 概念设计", 1)
    paragraph(doc, "系统按组织与培训基础、需求与报名、培训执行与成果三个模块进行概念设计。主业务链路为：维护基础数据 -> 发布课程并占用预算 -> 员工申请 -> 主管审批 -> HR 备案 -> 员工报名 -> 签到或缺勤 -> 完成培训 -> 评分、测试和证书。")
    heading(doc, "2.1 总体 E-R 图", 2)
    paragraph(doc, "总体模型由 13 张核心表组成。部门分别与员工、课程关联；员工通过角色表取得授权；课程设置唯一主讲师；申请成为报名的来源；报名是签到和培训完成的业务凭证。")
    doc.add_picture(str(ASSET_DIR / "overall.png"), width=Cm(16.2))
    caption(doc, "图 2-1 企业内部培训管理系统总体 E-R 图")
    heading(doc, "2.2 组织与培训基础模块 E-R 图", 2)
    paragraph(doc, "该模块维护部门、员工、角色、员工角色、黑名单、讲师和课程。员工归属部门，课程由部门主办；员工与角色为多对多关系；员工与黑名单、讲师与课程均为一对多关系。")
    doc.add_picture(str(ASSET_DIR / "basic.png"), width=Cm(16.2))
    caption(doc, "图 2-2 组织与培训基础模块 E-R 图")
    heading(doc, "2.3 需求与报名模块 E-R 图", 2)
    paragraph(doc, "员工对课程提交申请，申请经过部门审批和 HR 备案后才可产生报名。报名引用申请的 REQUEST_ID、EMP_ID、COURSE_ID 复合键，防止报名人与申请人或课程不一致。")
    doc.add_picture(str(ASSET_DIR / "request.png"), width=Cm(16.2))
    caption(doc, "图 2-3 需求与报名模块 E-R 图")
    heading(doc, "2.4 培训执行与成果模块 E-R 图", 2)
    paragraph(doc, "报名是签到的唯一来源，一条报名至多有一次签到。员工在完成培训后可进行讲师评分，HR 录入训前、训后测试并在满足条件后发放证书。")
    doc.add_picture(str(ASSET_DIR / "result.png"), width=Cm(16.2))
    caption(doc, "图 2-4 培训执行与成果模块 E-R 图")

    heading(doc, "3. 逻辑设计", 1)
    paragraph(doc, "数据库采用 Oracle identity 主键，所有业务时间使用 TIMESTAMP(0)，日历日期使用 DATE。表结构、字段名、约束与后端 Repository SQL 一致，数据库账户仅运行迁移和应用所需的最小权限。")
    heading(doc, "3.1 表的设计", 2)
    summary_rows = [(str(i), table, cname, purpose) for i, (table, cname, purpose, _) in enumerate(TABLES, 1)]
    add_table(doc, ["序号", "表名", "中文名称", "主要用途"], summary_rows, (1.2, 4.2, 3.2, 7.3))
    caption(doc, "表 3-1 核心业务表总表")
    for idx, (table, cname, purpose, fields) in enumerate(TABLES, 1):
        heading(doc, f"3.1.{idx} {table} 表", 3)
        paragraph(doc, purpose)
        add_table(doc, ["字段名", "数据类型", "约束", "字段说明"], fields)
        caption(doc, f"表 3-{idx + 1} {table} 表")

    heading(doc, "3.2 完整性、状态与索引设计", 2)
    paragraph(doc, "系统通过外键保证组织、申请、报名和成果记录的引用完整性，通过检查约束限制状态与数值范围，并以函数唯一索引防止并发场景下产生重复的有效业务记录。")
    rule_rows = [
        ("状态值", "员工 ACTIVE/RESIGNED；课程 DRAFT/PUBLISHED/CLOSED；申请 PENDING/DEPT_APPROVED/DEPT_REJECTED/HR_FILED；报名 REGISTERED/SIGNED_IN/COMPLETED/ABSENT/CANCELED。"),
        ("报名来源", "TRAINING_REGISTRATIONS 的 REQUEST_ID、EMP_ID、COURSE_ID 复合外键引用 TRAINING_REQUESTS，保证报名可追溯。"),
        ("唯一规则", "每员工最多一条 ACTIVE 黑名单；每员工每课程最多一条活跃申请和一条非取消报名；每类测试、评分和证书均有业务唯一约束。"),
        ("签到规则", "TRAINING_ATTENDANCE.REG_ID 唯一。一条报名仅允许一次签到；补签必须填写备注。"),
        ("迟到学时", "迟到分钟数取实际签到时间与开课时间的整分钟差；扣减学时按课程学时同比例计算，四舍五入至 0.5 小时，上限为课程学时。"),
        ("预算规则", "课程发布时校验部门余额并同步更新已使用预算；余额为虚拟列，不单独写入。"),
    ]
    add_table(doc, ["规则", "实现说明"], rule_rows, (3.2, 12.7))
    caption(doc, "表 3-15 关键完整性与业务规则")
    paragraph(doc, "索引包括员工部门状态、课程状态时间、申请状态课程、报名课程状态、评分讲师、测试课程和证书课程等查询索引；函数唯一索引作为 Service 层校验的最终数据库兜底。")

    heading(doc, "3.3 数据库关系图", 2)
    paragraph(doc, "关系图标出 13 张表的主要主键和外键关系。课程与讲师为一对多，培训申请与报名通过复合外键建立来源关系，报名与签到为一对零或一关系。")
    doc.add_picture(str(ASSET_DIR / "physical.png"), width=Cm(16.2))
    caption(doc, "图 3-1 数据库关系图")

    p = doc.add_paragraph(style="附录")
    run = p.add_run("附录A. 图表索引")
    set_east_asia(run, "黑体", 15, True)
    for text in ["图 2-1 企业内部培训管理系统总体 E-R 图", "图 2-2 组织与培训基础模块 E-R 图", "图 2-3 需求与报名模块 E-R 图", "图 2-4 培训执行与成果模块 E-R 图", "图 3-1 数据库关系图", "表 3-1 核心业务表总表", "表 3-2 至表 3-14 各实体字段设计表", "表 3-15 关键完整性与业务规则"]:
        paragraph(doc, text, first_line=False)


def main():
    if not TEMPLATE.exists():
        raise SystemExit(f"Template not found: {TEMPLATE}")
    build_diagrams()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    copyfile(TEMPLATE, OUTPUT)
    doc = Document(OUTPUT)
    remove_legacy_heading_numbering(doc)
    clear_after_cover(doc)
    add_toc(doc)
    add_content(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
